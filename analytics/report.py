import os
import pandas as pd
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_pdf_report(df_cleaned, profiling_res, kpis_res, anomalies_res, device_info, output_path="data/reports/SnapInsight_Report.pdf"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    
    # Custom Brand Styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=24,
        textColor=colors.HexColor('#0F172A'), # Slate 900
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        textColor=colors.HexColor('#2563EB'), # Blue 600
        spaceAfter=15
    )
    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6,
        leading=14
    )

    story = []

    # Title & Header
    story.append(Paragraph("SnapInsight — Executive Data Intelligence Report", title_style))
    story.append(Paragraph("Private On-Device AI Analytics Accelerated for Snapdragon® PCs", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#E2E8F0'), spaceAfter=15))

    # Dataset & Quality Overview
    story.append(Paragraph("1. Dataset & Quality Summary", section_heading))
    overview_text = (
        f"<b>Generated On:</b> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}<br/>"
        f"<b>Total Rows Analyzed:</b> {profiling_res.get('total_rows', len(df_cleaned)):,}<br/>"
        f"<b>Total Columns:</b> {profiling_res.get('total_cols', len(df_cleaned.columns))}<br/>"
        f"<b>Data Quality Score:</b> <font color='#16A34A'><b>{profiling_res.get('quality_score', 95)}/100</b></font><br/>"
        f"<b>Missing Values Cleaned:</b> {profiling_res.get('missing_values_count', 0)} | <b>Duplicates Removed:</b> {profiling_res.get('duplicate_rows_count', 0)}"
    )
    story.append(Paragraph(overview_text, body_style))
    story.append(Spacer(1, 10))

    # KPIs Section
    story.append(Paragraph("2. Key Performance Indicators", section_heading))
    kpis = kpis_res.get("kpis", {})
    kpi_data = [
        ["Total Revenue", "Avg Order Value", "Total Orders", "Growth Rate %", "Top Performing Product"],
        [
            f"₹{kpis.get('total_revenue', 0):,.2f}",
            f"₹{kpis.get('average_order_value', 0):,.2f}",
            f"{kpis.get('total_orders', 0):,}",
            f"{kpis.get('growth_rate_pct', 0)}%",
            str(kpis.get('top_product', 'N/A'))
        ]
    ]
    t_kpis = Table(kpi_data, colWidths=[100, 105, 80, 85, 160])
    t_kpis.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#0F172A')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_kpis)
    story.append(Spacer(1, 15))

    # Statistical Anomalies
    story.append(Paragraph("3. Statistical Anomaly Detection (IQR Method)", section_heading))
    anom_cnt = anomalies_res.get('anomalies_count', 0)
    target_col = anomalies_res.get('target_column', 'Revenue')
    anom_summary = (
        f"Analyzed feature <b>'{target_col}'</b> using Interquartile Range (IQR multiplier: {anomalies_res.get('iqr_multiplier', 1.5)}).<br/>"
        f"<b>Detected Anomalies:</b> {anom_cnt} records ({anomalies_res.get('anomalies_percentage', 0)}% of dataset).<br/>"
        f"<b>Normal Range Boundary:</b> Lower Bound: ₹{anomalies_res.get('lower_bound', 0):,.2f} | Upper Bound: ₹{anomalies_res.get('upper_bound', 0):,.2f}"
    )
    story.append(Paragraph(anom_summary, body_style))
    story.append(Spacer(1, 10))

    # Device & Runtime Info
    story.append(Paragraph("4. Hardware & AI Runtime Verification", section_heading))
    dev_text = (
        f"<b>Target Platform:</b> {device_info.get('device', 'Snapdragon-powered HP PC')}<br/>"
        f"<b>Processor:</b> {device_info.get('processor', 'Qualcomm Snapdragon X Elite / Host CPU')}<br/>"
        f"<b>Execution Mode:</b> <font color='#2563EB'><b>{device_info.get('execution', 'Snapdragon-Optimized / NPU-Ready (CPU Fallback)')}</b></font><br/>"
        f"<b>Local Privacy Guarantee:</b> All data processed locally on-device without cloud external transmission."
    )
    story.append(Paragraph(dev_text, body_style))
    story.append(Spacer(1, 15))

    doc.build(story)
    print(f"PDF Report successfully created at '{output_path}'.")
    return output_path


def generate_excel_report(df_cleaned, profiling_res, kpis_res, anomalies_res, device_info, output_path="data/reports/SnapInsight_Report.xlsx"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()

    # Styling constants
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # ----------------------------------------------------
    # Sheet 1: Summary
    # ----------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Summary"

    ws_summary.append(["SnapInsight — Data Intelligence Executive Summary"])
    ws_summary.cell(row=1, column=1).font = title_font
    ws_summary.append([])

    ws_summary.append(["Key Performance Indicator", "Value"])
    for cell in ws_summary[3]:
        cell.fill = header_fill
        cell.font = header_font

    kpis = kpis_res.get("kpis", {})
    summary_rows = [
        ("Total Revenue (₹)", kpis.get("total_revenue", 0.0)),
        ("Average Order Value (₹)", kpis.get("average_order_value", 0.0)),
        ("Total Orders Processed", kpis.get("total_orders", 0)),
        ("Period-over-Period Growth (%)", kpis.get("growth_rate_pct", 0.0)),
        ("Top Product", kpis.get("top_product", "N/A")),
        ("Data Quality Score (/100)", profiling_res.get("quality_score", 95)),
        ("Detected Anomalies Count", anomalies_res.get("anomalies_count", 0)),
        ("AI Execution Mode", device_info.get("execution", "Snapdragon-Optimized / NPU-Ready (CPU Fallback)"))
    ]

    for label, val in summary_rows:
        ws_summary.append([label, val])

    # ----------------------------------------------------
    # Sheet 2: Cleaned Data
    # ----------------------------------------------------
    ws_cleaned = wb.create_sheet(title="Cleaned Data")
    ws_cleaned.append(list(df_cleaned.columns))
    for cell in ws_cleaned[1]:
        cell.fill = header_fill
        cell.font = header_font

    for _, row in df_cleaned.head(5000).iterrows():
        ws_cleaned.append([None if pd.isnull(v) else v for v in row])

    # ----------------------------------------------------
    # Sheet 3: Anomalies
    # ----------------------------------------------------
    ws_anom = wb.create_sheet(title="Anomalies")
    ws_anom.append(["Row Index", "Target Column", "Value", "Anomaly Type", "Lower Bound", "Upper Bound"])
    for cell in ws_anom[1]:
        cell.fill = header_fill
        cell.font = header_font

    for anom in anomalies_res.get("anomalies", []):
        ws_anom.append([
            anom.get("row_index"),
            anom.get("target_column"),
            anom.get("value"),
            anom.get("anomaly_type"),
            anom.get("lower_bound"),
            anom.get("upper_bound")
        ])

    wb.save(output_path)
    print(f"Excel Report successfully created at '{output_path}'.")
    return output_path
